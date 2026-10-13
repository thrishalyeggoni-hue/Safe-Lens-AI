"""Isolated PDF/DOCX parser and inert PDF renderer. Never executes actions."""
import json
import logging
import math
from pathlib import Path
import re
import sys
import warnings
import zipfile

logging.disable(logging.CRITICAL)
warnings.filterwarnings('ignore')
MAX_PAGES=30
MAX_TEXT=12000
MAX_TOTAL_TEXT=60000
MAX_RENDER_PIXELS=4_000_000
MAX_RENDER_PAGES=3


def resolve(value):
    return value.get_object() if hasattr(value,'get_object') else value


def inspect_actions(obj,location):
    found=[]
    obj=resolve(obj)
    if not hasattr(obj,'get'): return found
    for key in ('/A','/OpenAction','/AA'):
        value=resolve(obj.get(key))
        if not hasattr(value,'get'): continue
        candidates=list(value.values())[:32] if key=='/AA' else [value]
        for action in candidates:
            action=resolve(action)
            if hasattr(action,'get') and str(action.get('/S')) in ('/JavaScript','/Launch','/GoToR','/SubmitForm','/ImportData'):
                found.append({'location':location,'action':str(action['/S']),'source':'pdf_structure','interpretation':'Observed active/external action; presence alone does not prove maliciousness. It was not executed.'})
    return found


def parse_pdf(path,folder):
    from pypdf import PdfReader, overwrite_configuration
    import pdfplumber
    import pypdfium2 as pdfium
    overwrite_configuration(maximum_declared_stream_length=32*1024*1024,zlib_maximum_output_length=32*1024*1024,lzw_maximum_output_length=32*1024*1024,run_length_maximum_output_length=32*1024*1024,stream_decoding_work_maximum_length=64*1024*1024,page_tree_maximum_entries=2000,xform_maximum_invocations_per_extraction=1000,disable_legacy_handling=True)
    reader=PdfReader(path,strict=True)
    result={'format':'PDF','parse_status':'completed','encrypted':bool(reader.is_encrypted),'metadata':{},'pages':[],'links':[],'structure_observations':[],'total_pages':None,'pages_skipped':[],'parse_failures':[],'parser':{'pypdf_strict':True,'native_text':'pdfplumber with pypdf fallback','renderer':'PDFium, forms not initialized, no JavaScript execution','parser_timeout_seconds':30,'render_max_pixels':MAX_RENDER_PIXELS}}
    if reader.is_encrypted:
        result.update(parse_status='encrypted',error='Encrypted PDF requires a password; no password guessing is attempted');return result
    result['total_pages']=len(reader.pages)
    result['metadata']={str(k):str(v)[:2000] for k,v in (reader.metadata or {}).items() if v is not None}
    root=resolve(reader.trailer['/Root']);result['structure_observations']+=inspect_actions(root,'document_catalog')
    names=resolve(root.get('/Names',{}))
    for key in ('/JavaScript','/EmbeddedFiles'):
        if key in names: result['structure_observations'].append({'location':'document_catalog','action':key,'source':'pdf_structure','interpretation':'Observed named scripts/attachments; not executed or extracted.'})
    count=min(MAX_PAGES,result['total_pages']); result['pages_skipped']=list(range(count+1,result['total_pages']+1))
    native=pdfplumber.open(path); renderer=None;rendered=0; total_chars=0
    try:
        for i in range(count):
            page=reader.pages[i]; number=i+1
            record={'page':number,'reference':f'Page {number}','source':'pdf_text','text':'','status':'completed','ocr_needed':False,'text_truncated':False,'words':[],'tables':[],'width_points':float(page.mediabox.width),'height_points':float(page.mediabox.height)}
            result['structure_observations']+=inspect_actions(page,number)
            try:
                pl=native.pages[i]
                text=pl.extract_text() or page.extract_text() or ''
                allowed=max(0,min(MAX_TEXT,MAX_TOTAL_TEXT-total_chars));record['text']=text[:allowed];record['text_truncated']=len(text)>allowed;total_chars+=len(record['text'])
                record['words']=[{k:w[k] for k in ('text','x0','top','x1','bottom')} for w in pl.extract_words()[:3000]]
                record['tables']=[[[str(cell)[:500] if cell is not None else None for cell in row[:20]] for row in table[:50]] for table in pl.extract_tables()[:10]]
                record['ocr_needed']=sum(c.isalnum() for c in record['text'])<12 and not record['text_truncated']
                if record['ocr_needed']: record['source']='ocr_pending'
            except Exception as exc:
                record.update(status='failed',error=f'Page extraction failed ({type(exc).__name__})',ocr_needed=True,source='ocr_pending');result['parse_failures'].append(number)
            for ann in resolve(page.get('/Annots',[]))[:200]:
                obj=resolve(ann);result['structure_observations']+=inspect_actions(obj,number)
                action=resolve(obj.get('/A',{})); uri=action.get('/URI') if hasattr(action,'get') else None
                if uri is not None:
                    rect=[float(x) for x in resolve(obj.get('/Rect',[]))]
                    displayed=[]
                    if len(rect)==4 and not page.rotation and page.cropbox==page.mediabox:
                        x0,y0,x1,y1=rect;top=record['height_points']-max(y0,y1);bottom=record['height_points']-min(y0,y1)
                        displayed=[w['text'] for w in record['words'] if w['x1']>min(x0,x1) and w['x0']<max(x0,x1) and w['bottom']>top and w['top']<bottom]
                    result['links'].append({'page':number,'source':'pdf_embedded_hyperlink','url':str(uri)[:12000],'rectangle_points':rect,'displayed_text':' '.join(displayed),'destination_truncated':len(str(uri))>12000})
            # Render first page for preview, plus bounded scanned pages for OCR.
            if (i==0 or record['ocr_needed']) and rendered<MAX_RENDER_PAGES:
                try:
                    if renderer is None: renderer=pdfium.PdfDocument(path)
                    rp=renderer[i];width,height=rp.get_size()
                    if not math.isfinite(width*height) or min(width,height)<=0 or max(width,height)>14400: raise ValueError('Unsafe page dimensions')
                    scale=min(2.0,1600/max(width,height),math.sqrt(MAX_RENDER_PIXELS/(width*height)))
                    bitmap=rp.render(scale=scale,may_draw_forms=False); image=bitmap.to_pil();dest=Path(folder)/f'page_{number}.png';image.save(dest)
                    record['rendered_image']=dest.name; record['render_width']=image.width;record['render_height']=image.height;record['render_scale']=scale;record['coordinate_space']='rendered page pixels after PDFium page rotation; PDF annotations use original PDF point coordinates'
                    image.close();bitmap.close();rp.close();rendered+=1
                except Exception as exc: record['render_error']=f'Page rendering unavailable ({type(exc).__name__})'
            elif record['ocr_needed']: record['render_error']='Scanned-page render/OCR limit reached'
            result['pages'].append(record)
    finally:
        native.close()
        if renderer is not None: renderer.close()
    return result


def parse_docx(path,folder):
    from docx import Document
    document=Document(path)
    result={'format':'DOCX','parse_status':'completed','encrypted':False,'metadata':{'title':document.core_properties.title,'author':document.core_properties.author,'created':str(document.core_properties.created) if document.core_properties.created else None,'modified':str(document.core_properties.modified) if document.core_properties.modified else None},'pages':[],'links':[],'structure_observations':[],'total_pages':None,'pages_skipped':[],'parse_failures':[],'pagination_limitation':'DOCX has no fixed parsed page numbers. References identify paragraphs/tables, not invented pages.','parser':{'engine':'python-docx','logical_block_limit':100,'total_text_limit':MAX_TOTAL_TEXT}}
    from docx.oxml.ns import qn
    total=0
    for i,paragraph in enumerate(document.paragraphs[:100],1):
        text=paragraph.text;allowed=max(0,min(MAX_TEXT,MAX_TOTAL_TEXT-total));total+=min(len(text),allowed)
        result['pages'].append({'page':None,'reference':f'Paragraph {i}','source':'docx_text','text':text[:allowed],'status':'completed','ocr_needed':False,'text_truncated':len(text)>allowed})
        for hyperlink in paragraph._p.iter(qn('w:hyperlink')):
            rid=hyperlink.get(qn('r:id')); rel=document.part.rels.get(rid)
            if rel and rel.is_external: result['links'].append({'page':None,'reference':f'Paragraph {i}','source':'docx_embedded_hyperlink','url':rel.target_ref,'displayed_text':''.join(node.text or '' for node in hyperlink.iter(qn('w:t')))})
    for i,table in enumerate(document.tables[:20],1):
        text='\n'.join(' | '.join(cell.text for cell in row.cells[:20]) for row in table.rows[:50]);allowed=max(0,min(MAX_TEXT,MAX_TOTAL_TEXT-total));total+=min(len(text),allowed)
        result['pages'].append({'page':None,'reference':f'Table {i}','source':'docx_table','text':text[:allowed],'status':'completed','ocr_needed':False,'text_truncated':len(text)>allowed})
    if len(document.paragraphs)>100 or len(document.tables)>20: result['blocks_skipped']=True
    result['unprocessed_parts']=[]
    with zipfile.ZipFile(path) as zipped:
        for name in zipped.namelist():
            if name.startswith(('word/media/','word/header','word/footer','word/footnotes','word/endnotes','word/comments')): result['unprocessed_parts'].append(name[:200])
            if name.startswith(('word/embeddings/','word/activeX/')): result['structure_observations'].append({'location':name[:200],'action':'Embedded object','source':'docx_structure','interpretation':'Embedded object not extracted or executed.'})
    return result


if __name__=='__main__':
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_CPU,(25,25))
        memory_applied=True
        try: resource.setrlimit(resource.RLIMIT_AS,(1536*1024*1024,1536*1024*1024))
        except (ValueError,OSError): memory_applied=False
        path,folder,fmt=sys.argv[1:4]
        result=parse_pdf(path,folder) if fmt=='PDF' else parse_docx(path,folder)
        result['resource_limits']={'cpu_seconds':25,'address_space_bytes':1536*1024*1024,'address_space_applied':memory_applied}
    except Exception as exc: result={'parse_status':'failed','pages':[],'links':[],'total_pages':None,'encrypted':False,'metadata':{},'structure_observations':[],'parse_failures':[],'pages_skipped':[],'error':f'Document parsing failed ({type(exc).__name__})'}
    Path(sys.argv[2],'parsed.json').write_text(json.dumps(result,ensure_ascii=False))
