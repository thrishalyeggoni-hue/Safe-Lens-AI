import asyncio
import io
from pathlib import Path
from unittest.mock import AsyncMock
import zipfile
import pytest
from reportlab.pdfgen import canvas
from pypdf import PdfReader,PdfWriter
from backend.app.services import document_service as service
from backend.app.config import settings
from tests.test_link_api import client


def pdf_bytes(lines=None,pages=1):
    out=io.BytesIO();c=canvas.Canvas(out,pageCompression=0)
    for _ in range(pages):
        for i,line in enumerate(lines or []): c.drawString(30,730-i*25,line)
        c.showPage()
    c.save();return out.getvalue()


def clean_scan(monkeypatch):
    monkeypatch.setattr(service,'scan_document',AsyncMock(return_value={'status':'clean','detail':'Unit test scanner stub, not real scan','signature_match':None}))
    monkeypatch.setattr(settings,'MESSAGE_AI_ENABLED',False)


def test_contents_not_name():
    assert service.validate_document(pdf_bytes(['hello']),'wrong.exe')['validated_type']=='PDF'
    from docx import Document
    out=io.BytesIO();doc=Document();doc.add_paragraph('Controlled document');doc.save(out)
    assert service.validate_document(out.getvalue(),'wrong.pdf')['validated_type']=='DOCX'

@pytest.mark.parametrize('raw',[b'',b'not a document',b'PK\x03\x04bad',b'%PDF-1.4\n'+b'x'*(service.MAX_BYTES+1)])
def test_invalid_contents_and_limits(raw):
    with pytest.raises(ValueError): service.validate_document(raw,'test.pdf')


def test_zip_bomb_rejected():
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml','wordprocessingml.document.main+xml');z.writestr('word/document.xml','x'*1000000)
    with pytest.raises(ValueError,match='expansion'): service.validate_document(out.getvalue(),'test.docx')


@pytest.mark.asyncio
async def test_real_text_parser_exact_page_evidence(monkeypatch):
    clean_scan(monkeypatch)
    result=await service.analyze_document_bytes(pdf_bytes(['Pay INR 2500 registration fee today for your job.','Send me your OTP now.']),'job.pdf')
    assert result['document']['total_pages']==1 and result['document']['parse_status']=='completed'
    assert 'Send me your OTP now.' in result['pages'][0]['text']
    assert result['risk_category']=='High Risk' and result['score'] is None
    assert all(f['page_number']==1 for f in result['grounded_findings'])
    for f in result['grounded_findings']: assert result['pages'][0]['text'][f['start']:f['end']]==f['text']
    assert result['coverage']['pages_ocr_processed']==0


@pytest.mark.asyncio
async def test_encrypted_and_corrupt_honest(monkeypatch):
    clean_scan(monkeypatch)
    writer=PdfWriter();writer.append(PdfReader(io.BytesIO(pdf_bytes(['secret']))));writer.encrypt('controlled');out=io.BytesIO();writer.write(out)
    for raw in [out.getvalue(),b'%PDF-1.7\ncorrupt']:
        result=await service.analyze_document_bytes(raw,'test.pdf')
        assert result['risk_category']=='Unable to Assess' and result['pages']==[] and not result['coverage']['analysis_complete']
    assert result['document']['parse_status']=='failed'


@pytest.mark.asyncio
async def test_scanned_page_with_unavailable_ocr_not_safe(monkeypatch):
    clean_scan(monkeypatch)
    monkeypatch.setattr(service,'run_paddle',AsyncMock(return_value={'status':'unavailable','regions':[],'runs':[],'error':'Unit test missing OCR'}))
    result=await service.analyze_document_bytes(pdf_bytes(),'blank.pdf')
    assert result['coverage']['pages_ocr_processed']==1
    assert result['pages'][0]['source']=='ocr' and result['pages'][0]['status']=='unreadable'
    assert result['risk_category']=='Unable to Assess'


@pytest.mark.asyncio
async def test_limit_and_no_silent_text_truncation(monkeypatch):
    clean_scan(monkeypatch)
    result=await service.analyze_document_bytes(pdf_bytes(['Normal controlled text document.'],pages=32),'many.pdf')
    assert result['coverage']['total_pages']==32 and result['coverage']['pages_parsed']==30
    assert result['coverage']['pages_skipped']==[31,32] and not result['coverage']['analysis_complete']
    assert result['risk_category']!='Low Risk'


@pytest.mark.asyncio
async def test_malware_match_prevents_parser(monkeypatch):
    monkeypatch.setattr(service,'scan_document',AsyncMock(return_value={'status':'infected','signature_match':'Unit-test signature','detail':'Explicit test stub'}))
    parser=AsyncMock();monkeypatch.setattr(service,'parse_real',parser)
    result=await service.analyze_document_bytes(pdf_bytes(['ordinary']),'test.pdf')
    assert result['risk_category']=='High Risk' and result['malware']['status']=='infected'
    parser.assert_not_called(); assert result['pages']==[]


@pytest.mark.asyncio
async def test_native_link_mismatch_reuses_real_checker_interface(monkeypatch):
    clean_scan(monkeypatch)
    checker=AsyncMock(return_value={'risk_category':'Unable to Assess','findings':[],'checks_performed':[],'unavailable_checks':[]})
    monkeypatch.setattr(service,'analyze_link',checker)
    out=io.BytesIO();c=canvas.Canvas(out);c.drawString(40,700,'https://example.com');c.linkURL('https://example.org',(38,695,350,718));c.save()
    result=await service.analyze_document_bytes(out.getvalue(),'mismatch.pdf')
    assert {item['source'] for item in result['url_assessments']}=={'pdf_embedded_hyperlink','visible_url_text'}
    assert any(f['source']=='PDF/DOCX hyperlink structure' and f['page_number']==1 for f in result['findings'])
    assert result['risk_category']=='Caution' and checker.await_count==2


def test_document_api_privacy_stream_and_session(client,monkeypatch):
    clean_scan(monkeypatch)
    response=client.post('/api/check/document?stream=true',files={'file':('test.pdf',pdf_bytes(['Send me your OTP now. Private bank details 123456789.']),'application/pdf')})
    import json
    events=[json.loads(line) for line in response.text.splitlines()];result=events[-1]['data']
    assert result['risk_category']=='High Risk' and response.headers['cache-control']=='no-store'
    assert any('Parsing' in e.get('stage','') for e in events)
    assert '123456789' not in client.get('/api/history').text
    assert client.get('/api/assessments/'+result['id']).json()['document_details_retained'] is False
    client.cookies.clear();assert client.get('/api/assessments/'+result['id']).status_code==403


def test_document_upload_rejection_before_parser(client):
    assert client.post('/api/check/document',files={'file':('fake.pdf',b'not a pdf','application/pdf')}).status_code==422
    assert client.post('/api/check/document',content=b'x',headers={'Content-Length':str(22*1024*1024),'Content-Type':'multipart/form-data; boundary=test'}).status_code==413


@pytest.mark.asyncio
async def test_parser_cancellation_kills_worker(monkeypatch,tmp_path):
    started=asyncio.Event();process=None
    class Process:
        returncode=None;killed=False
        async def wait(self):
            if self.killed: self.returncode=-9;return -9
            started.set();await asyncio.Event().wait()
        def kill(self):self.killed=True
    async def spawn(*args,**kwargs):
        nonlocal process
        process=Process();return process
    monkeypatch.setattr(asyncio,'create_subprocess_exec',spawn)
    task=asyncio.create_task(service.parse_real(tmp_path/'test.pdf',tmp_path,'PDF'));await started.wait();task.cancel()
    with pytest.raises(asyncio.CancelledError):await task
    assert process.killed


@pytest.mark.asyncio
async def test_docx_real_parser_no_invented_pages(monkeypatch):
    clean_scan(monkeypatch)
    from docx import Document
    d=Document();d.add_paragraph('Send me your OTP now.');out=io.BytesIO();d.save(out)
    result=await service.analyze_document_bytes(out.getvalue(),'test.docx')
    assert result['document']['total_pages'] is None and result['pages'][0]['reference']=='Paragraph 1'
    assert result['grounded_findings'][0]['page_number'] is None and result['risk_category']=='High Risk'

@pytest.mark.asyncio
async def test_uncertain_ocr_cannot_assert_direct_high_risk_or_contact_urls(monkeypatch):
    clean_scan(monkeypatch)
    text='Send me your OTP now. Visit https://example.com'
    region={'text':text,'confidence':.3,'bounding_box':{'points':[[5,5],[500,5],[500,40],[5,40]]},'language_model':'en','model':'explicit-test-stub'}
    monkeypatch.setattr(service,'run_paddle',AsyncMock(return_value={'status':'completed','regions':[region],'runs':[]}))
    checker=AsyncMock();monkeypatch.setattr(service,'analyze_link',checker)
    result=await service.analyze_document_bytes(pdf_bytes(),'test.pdf')
    assert result['risk_category']=='Caution' and not result['coverage']['analysis_complete']
    assert result['grounded_findings'][0]['ocr_certainty']=='uncertain'
    checker.assert_not_called()


@pytest.mark.asyncio
async def test_docx_unprocessed_images_disclose_missing_coverage(monkeypatch):
    clean_scan(monkeypatch)
    from docx import Document
    from PIL import Image
    image=io.BytesIO();Image.new('RGB',(20,20),'white').save(image,format='PNG');image.seek(0)
    d=Document();d.add_paragraph('Controlled ordinary meeting agenda.');d.add_picture(image);out=io.BytesIO();d.save(out)
    result=await service.analyze_document_bytes(out.getvalue(),'image.docx')
    assert result['coverage']['unprocessed_parts'] and not result['coverage']['parse_coverage_complete']
    assert result['risk_category']!='Low Risk'


@pytest.mark.asyncio
@pytest.mark.parametrize('page,quote,accepted',[(2,'Send me your OTP',True),(9,'Send me your OTP',False),(2,'invented evidence',False)])
async def test_document_qwen_requires_actual_page_and_quote(monkeypatch,page,quote,accepted):
    from backend.app.services import message_ai
    from backend.app.services.message_extraction import extract_message
    import json
    requests=[]
    async def read(client,method,url,**kwargs):
        if method=='GET':return {'models':[{'name':settings.MESSAGE_OLLAMA_MODEL}]}
        requests.append(kwargs['json'])
        concern={'category':'credential_request','evidence':quote,'interpretation':'Observed request, sender unverified.','recommended_action':'Verify independently.','uncertainty':'Context uncertain.','page_number':page,'reference':'Page 2'}
        return {'done':True,'model':settings.MESSAGE_OLLAMA_MODEL,'message':{'content':json.dumps({'context':'direct_request','concerns':[concern]})}}
    monkeypatch.setattr(settings,'MESSAGE_AI_ENABLED',True);monkeypatch.setattr(message_ai,'read_json',read)
    text='Send me your OTP'
    result=await message_ai.interpret_message(text,extract_message(text),[],document_context={'page_number':2,'reference':'Page 2'})
    assert bool(result['interpretations'])==accepted
    if accepted:assert result['interpretations'][0]['page_number']==2
    else:assert result['grounding_failures']
    assert json.loads(requests[0]['messages'][1]['content'])['document_context']['page_number']==2
    assert 'page_number' in requests[0]['format']['$defs']['DocumentConcern']['required']
